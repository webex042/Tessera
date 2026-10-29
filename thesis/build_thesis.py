#!/usr/bin/env python3
"""Build the BIT Mesra B.Tech thesis DOCX for the Tessera project.

Usage:  python3 build_thesis.py [output.docx]

Reads the markdown-like sources in src/, the table data in src/tables.py and the
figure registry in src/figures.py, and emits a single formatted DOCX.
"""
import re
import sys
import pathlib

from docx import Document
from docx.shared import Pt, Inches, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = pathlib.Path(__file__).resolve().parent
SRC = HERE / "src"
FIGDIR = HERE / "figures"
sys.path.insert(0, str(SRC))
from tables import TABLES          # noqa: E402
from figures import FIGURES        # noqa: E402

BODY_FONT = "Times New Roman"
MONO_FONT = "Consolas"
BODY_SIZE = Pt(12)

# ----------------------------------------------------------------- details
STUDENT = "Pranshu Ranjan"
ROLL = "BTECH/15191/24"
DEGREE = "Bachelor of Technology"
BRANCH = "Computer Science and Engineering"
GUIDE = "Dr. Ayan Kumar Das"
TITLE = ("Tessera: A Privacy-Preserving, Hardware-Independent Local Inference "
         "and Screen-Aware Memory System for Windows")

# ----------------------------------------------------------------- helpers

def set_cell_background(cell, hex_colour):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_colour)
    tcPr.append(shd)


def shade_paragraph(par, hex_colour):
    pPr = par._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_colour)
    pPr.append(shd)


def add_field(par, instr):
    """Insert a Word field (used for TOC, page numbers, lists of figures)."""
    r = par.add_run()
    fld = OxmlElement("w:fldChar"); fld.set(qn("w:fldCharType"), "begin")
    r._r.append(fld)
    r2 = par.add_run()
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve")
    it.text = instr
    r2._r.append(it)
    r3 = par.add_run()
    sep = OxmlElement("w:fldChar"); sep.set(qn("w:fldCharType"), "separate")
    r3._r.append(sep)
    r4 = par.add_run("(update field to populate)")
    r4.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
    r5 = par.add_run()
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    r5._r.append(end)


def section_page_numbering(section, fmt, start=None):
    """fmt: 'lowerRoman' or 'decimal'."""
    sectPr = section._sectPr
    for old in sectPr.findall(qn("w:pgNumType")):
        sectPr.remove(old)
    pg = OxmlElement("w:pgNumType")
    pg.set(qn("w:fmt"), fmt)
    if start is not None:
        pg.set(qn("w:start"), str(start))
    sectPr.append(pg)


def footer_page_number(section, link_previous=False):
    footer = section.footer
    footer.is_linked_to_previous = link_previous
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.text = ""
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_field(p, " PAGE ")
    for r in p.runs:
        r.font.name = BODY_FONT
        r.font.size = Pt(11)
        r.font.color.rgb = RGBColor(0, 0, 0)
    # remove the grey placeholder run text that add_field inserts
    for r in p.runs:
        if r.text == "(update field to populate)":
            r.text = "1"


def enable_update_fields(doc):
    settings = doc.settings.element
    for old in settings.findall(qn("w:updateFields")):
        settings.remove(old)
    el = OxmlElement("w:updateFields")
    el.set(qn("w:val"), "true")
    settings.append(el)


def keep_with_next(par):
    pPr = par._p.get_or_add_pPr()
    k = OxmlElement("w:keepNext"); k.set(qn("w:val"), "true")
    pPr.append(k)


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader"); el.set(qn("w:val"), "true")
    trPr.append(el)



def add_cached_toc(doc, instr, entries, style_for_level):
    """Insert a real Word TOC field whose cached result is `entries`.

    entries: list of (level:int, text:str, page:str). Word repopulates the field on
    update; LibreOffice and any PDF export render the cached result as written.
    """
    if not entries:
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        add_field(p, instr)
        return

    paras = []
    for level, text, page in entries:
        p = doc.add_paragraph(style=style_for_level(level))
        pf = p.paragraph_format
        pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
        pf.space_after = Pt(3)
        pf.left_indent = Inches(0.0 if level <= 1 else 0.25 * (level - 1))
        _add_right_tab(p)
        r = p.add_run(text)
        r.font.name = BODY_FONT
        r.font.size = Pt(11.5 if level <= 1 else 11)
        r.bold = (level <= 1)
        r2 = p.add_run("\t" + page)
        r2.font.name = BODY_FONT
        r2.font.size = Pt(11.5 if level <= 1 else 11)
        r2.bold = (level <= 1)
        paras.append(p)

    first, last = paras[0], paras[-1]
    # begin + instruction + separate, prepended to the first result paragraph
    begin = OxmlElement("w:r")
    fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), "begin")
    begin.append(fc)
    instr_r = OxmlElement("w:r")
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = instr
    instr_r.append(it)
    sep = OxmlElement("w:r")
    sc = OxmlElement("w:fldChar"); sc.set(qn("w:fldCharType"), "separate")
    sep.append(sc)
    anchor = first._p
    for el in (sep, instr_r, begin):
        anchor.insert(1 if anchor.find(qn("w:pPr")) is not None else 0, el)
    end = OxmlElement("w:r")
    ec = OxmlElement("w:fldChar"); ec.set(qn("w:fldCharType"), "end")
    end.append(ec)
    last._p.append(end)


def _add_right_tab(par):
    pPr = par._p.get_or_add_pPr()
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "right")
    tab.set(qn("w:leader"), "dot")
    tab.set(qn("w:pos"), "9000")
    tabs.append(tab)
    pPr.append(tabs)

# ----------------------------------------------------------------- styles

def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = BODY_SIZE
    if normal.element.rPr is not None and normal.element.rPr.rFonts is not None:
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    pf = normal.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_after = Pt(8)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    h1 = doc.styles["Heading 1"]
    h1.font.name = BODY_FONT; h1.font.size = Pt(16); h1.font.bold = True
    h1.font.color.rgb = RGBColor(0, 0, 0)
    h1.paragraph_format.space_before = Pt(0)
    h1.paragraph_format.space_after = Pt(18)
    h1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h1.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

    h2 = doc.styles["Heading 2"]
    h2.font.name = BODY_FONT; h2.font.size = Pt(13.5); h2.font.bold = True
    h2.font.color.rgb = RGBColor(0, 0, 0)
    h2.paragraph_format.space_before = Pt(16)
    h2.paragraph_format.space_after = Pt(6)
    h2.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    h2.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

    h3 = doc.styles["Heading 3"]
    h3.font.name = BODY_FONT; h3.font.size = Pt(12); h3.font.bold = True
    h3.font.italic = True
    h3.font.color.rgb = RGBColor(0, 0, 0)
    h3.paragraph_format.space_before = Pt(12)
    h3.paragraph_format.space_after = Pt(4)
    h3.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    h3.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

    from docx.enum.style import WD_STYLE_TYPE
    for name, align in (("FigCaption", WD_ALIGN_PARAGRAPH.CENTER),
                        ("TblCaption", WD_ALIGN_PARAGRAPH.LEFT)):
        st = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        st.base_style = doc.styles["Normal"]
        st.font.name = BODY_FONT
        st.font.size = Pt(10.5)
        st.quick_style = False
        pfc = st.paragraph_format
        pfc.alignment = align
        pfc.line_spacing_rule = WD_LINE_SPACING.SINGLE
        pfc.space_after = Pt(14)

    for name in ("List Bullet", "List Number"):
        st = doc.styles[name]
        st.font.name = BODY_FONT
        st.font.size = BODY_SIZE
        st.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        st.paragraph_format.space_after = Pt(4)


# ----------------------------------------------------------------- inline

BOLD_RE = re.compile(r"\*\*(.+?)\*\*")
ITAL_RE = re.compile(r"(?<!\*)\*(?!\*)(.+?)(?<!\*)\*(?!\*)")
CODE_RE = re.compile(r"`([^`]+)`")


def add_rich(par, text, size=None, bold_all=False):
    """Add text with **bold**, *italic* and `code` spans."""
    token = re.compile(r"(\*\*.+?\*\*|(?<!\*)\*(?!\*).+?(?<!\*)\*(?!\*)|`[^`]+`)")
    for piece in token.split(text):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**"):
            r = par.add_run(piece[2:-2]); r.bold = True
        elif piece.startswith("`") and piece.endswith("`"):
            r = par.add_run(piece[1:-1]); r.font.name = MONO_FONT
            r.font.size = Pt((size or BODY_SIZE).pt - 1.5)
        elif piece.startswith("*") and piece.endswith("*") and len(piece) > 2:
            r = par.add_run(piece[1:-1]); r.italic = True
        else:
            r = par.add_run(piece)
        r.font.name = r.font.name or BODY_FONT
        if size:
            r.font.size = size
        if bold_all:
            r.bold = True
    return par


# ----------------------------------------------------------------- blocks

def add_figure(doc, fid):
    fname, caption, width = FIGURES[fid]
    path = FIGDIR / fname
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    p.add_run().add_picture(str(path), width=Inches(width))
    keep_with_next(p)
    c = doc.add_paragraph(style="FigCaption")
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.paragraph_format.space_after = Pt(14)
    c.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    r = c.add_run(f"Figure {fid}:  ")
    r.bold = True; r.font.size = Pt(10.5); r.font.name = BODY_FONT
    r2 = c.add_run(caption)
    r2.font.size = Pt(10.5); r2.font.name = BODY_FONT
    return fid, caption


def add_table(doc, tid):
    spec = TABLES[tid]
    cap = doc.add_paragraph(style="TblCaption")
    cap.alignment = WD_ALIGN_PARAGRAPH.LEFT
    cap.paragraph_format.space_before = Pt(12)
    cap.paragraph_format.space_after = Pt(4)
    cap.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    r = cap.add_run(f"Table {tid}:  ")
    r.bold = True; r.font.size = Pt(10.5); r.font.name = BODY_FONT
    r2 = cap.add_run(spec["caption"])
    r2.font.size = Pt(10.5); r2.font.name = BODY_FONT
    keep_with_next(cap)

    rows = spec["rows"]
    ncols = len(rows[0])
    t = doc.add_table(rows=0, cols=ncols)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    size = Pt(9.5) if ncols >= 4 else Pt(10)

    # Proportional column widths from the longest cell in each column, so short
    # identifier columns do not take the same space as prose columns.
    avail = 6.0  # inches of text width
    weights = []
    for j in range(ncols):
        longest = max(len(str(r[j])) for r in rows)
        weights.append(max(6.0, min(float(longest), 70.0)))
    total = sum(weights)
    widths = [Inches(max(0.55, avail * w / total)) for w in weights]
    scale = avail / sum(w.inches for w in widths)
    widths = [Inches(w.inches * scale) for w in widths]
    t.autofit = False
    tblPr = t._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    for i, row in enumerate(rows):
        cells = t.add_row().cells
        if i == 0:
            repeat_header(t.rows[0])
        for j, val in enumerate(row):
            cell = cells[j]
            cell.width = widths[j]
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            add_rich(p, val, size=size, bold_all=(i == 0))
            if i == 0:
                set_cell_background(cell, "E8E8E8")
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(10)
    after.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    return tid, spec["caption"]


def add_code(doc, lines):
    for i, line in enumerate(lines):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        p.paragraph_format.space_before = Pt(6) if i == 0 else Pt(0)
        p.paragraph_format.space_after = Pt(6) if i == len(lines) - 1 else Pt(0)
        p.paragraph_format.left_indent = Inches(0.25)
        shade_paragraph(p, "F4F4F4")
        r = p.add_run(line if line else " ")
        r.font.name = MONO_FONT
        r.font.size = Pt(9)


def add_placeholder(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.left_indent = Inches(0.2)
    p.paragraph_format.right_indent = Inches(0.2)
    shade_paragraph(p, "FFF2CC")
    r = p.add_run("[TO BE CONFIRMED] ")
    r.bold = True; r.font.name = BODY_FONT; r.font.size = Pt(11)
    add_rich(p, text, size=Pt(11))


# ----------------------------------------------------------------- parser

def render_source(doc, path, fig_list, tbl_list, outline=None, state=None):
    state = {} if state is None else state
    lines = pathlib.Path(path).read_text(encoding="utf-8").split("\n")
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip():
            i += 1
            continue

        if line.startswith("```"):
            block = []
            i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                block.append(lines[i].rstrip())
                i += 1
            i += 1
            add_code(doc, block)
            continue

        if line.startswith("### "):
            body = line[4:]
            num, _, text = body.partition("|")
            p = doc.add_heading(level=3)
            _t = f"{num}  {text}" if text else num
            if outline is not None:
                outline.append((3, re.sub(r"[*`]", "", _t)))
            add_rich(p, _t)
            for r in p.runs:
                r.font.name = BODY_FONT; r.font.size = Pt(12)
                r.bold = True; r.italic = True
                r.font.color.rgb = RGBColor(0, 0, 0)
            i += 1
            continue

        if line.startswith("## "):
            body = line[3:]
            num, _, text = body.partition("|")
            p = doc.add_heading(level=2)
            _t = f"{num}  {text}" if text else num
            if outline is not None:
                outline.append((2, re.sub(r"[*`]", "", _t)))
            add_rich(p, _t)
            for r in p.runs:
                r.font.name = BODY_FONT; r.font.size = Pt(13.5)
                r.bold = True; r.font.color.rgb = RGBColor(0, 0, 0)
            i += 1
            continue

        if line.startswith("# "):
            body = line[2:]
            label, _, text = body.partition("|")
            p = doc.add_heading(level=1)
            if not state.get("first_chapter"):
                p.paragraph_format.page_break_before = True
            state["first_chapter"] = False
            full = f"{label.strip()}  \u2014  {text.strip()}" if text.strip() else label.strip()
            if outline is not None:
                outline.append((1, full))
            r = p.add_run(full)
            r.font.name = BODY_FONT; r.font.size = Pt(15); r.bold = True
            r.font.color.rgb = RGBColor(0, 0, 0)
            i += 1
            continue

        if line.startswith("FIG:"):
            fig_list.append(add_figure(doc, line[4:].strip()))
            i += 1
            continue

        if line.startswith("TBL:"):
            tbl_list.append(add_table(doc, line[4:].strip()))
            i += 1
            continue

        if line.startswith("!!"):
            add_placeholder(doc, line[2:].strip())
            i += 1
            continue

        if line.startswith("@REF|"):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.left_indent = Inches(0.45)
            p.paragraph_format.first_line_indent = Inches(-0.45)
            p.paragraph_format.space_after = Pt(8)
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            add_rich(p, line[5:].strip(), size=Pt(11))
            i += 1
            continue

        if line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            add_rich(p, line[2:])
            i += 1
            continue

        if line.startswith("+ "):
            p = doc.add_paragraph(style="List Number")
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            add_rich(p, line[2:])
            i += 1
            continue

        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_rich(p, line)
        i += 1


# ----------------------------------------------------------------- prelims

def centred(doc, text, size=12, bold=False, italic=False, space_after=6,
            caps=False, underline=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    r = p.add_run(text.upper() if caps else text)
    r.font.name = BODY_FONT; r.font.size = Pt(size)
    r.bold = bold; r.italic = italic; r.underline = underline
    return p


def body_par(doc, text, size=12, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=10,
             spacing=WD_LINE_SPACING.ONE_POINT_FIVE, indent=0.0):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing_rule = spacing
    if indent:
        p.paragraph_format.left_indent = Inches(indent)
        p.paragraph_format.right_indent = Inches(indent)
    add_rich(p, text, size=Pt(size))
    return p


def build_title_page(doc):
    for _ in range(2):
        doc.add_paragraph()
    centred(doc, TITLE, size=16, bold=True, space_after=28)
    centred(doc, "A Thesis", size=13, space_after=4)
    centred(doc, "Submitted in partial fulfillment of the requirements", size=12, space_after=2)
    centred(doc, "for the award of the Degree of", size=12, space_after=14)
    centred(doc, DEGREE.upper(), size=13.5, bold=True, space_after=6)
    centred(doc, "IN", size=12, space_after=6)
    centred(doc, BRANCH.upper(), size=13.5, bold=True, space_after=24)
    centred(doc, "BY", size=12, space_after=8)
    centred(doc, STUDENT.upper(), size=13.5, bold=True, space_after=4)
    centred(doc, f"({ROLL})", size=12, space_after=24)
    centred(doc, "Under the supervision of", size=12, italic=True, space_after=6)
    centred(doc, GUIDE, size=13, bold=True, space_after=30)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("_" * 56)
    r.font.name = BODY_FONT
    centred(doc, "DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING", size=12, bold=True, space_after=4)
    centred(doc, "BIRLA INSTITUTE OF TECHNOLOGY", size=14, bold=True, space_after=4)
    centred(doc, "MESRA-835215, RANCHI", size=13, bold=True, space_after=4)
    centred(doc, "OFF-CAMPUS PATNA", size=12, bold=True, space_after=14)
    add_placeholder(doc, "Academic session and month/year of submission are not stated anywhere in "
                         "the project material. Please supply them (for example \"2026-27\" and "
                         "\"November 2026\") so they can be added beneath the institute address, and "
                         "confirm that the main-campus address above together with the off-campus "
                         "line is the arrangement your department requires.")


def build_approval(doc):
    doc.add_page_break()
    centred(doc, "APPROVAL OF THE GUIDE", size=14, bold=True, space_after=24)
    body_par(doc, "Recommended that the thesis entitled")
    body_par(doc, f"“{TITLE}”", align=WD_ALIGN_PARAGRAPH.CENTER)
    body_par(doc, f"presented by **{STUDENT}** ({ROLL}) under my supervision and guidance be "
                  f"accepted as fulfilling this part of the requirements for the award of the "
                  f"Degree of **{DEGREE} in {BRANCH}**. To the best of my knowledge, the content "
                  f"of this thesis did not form a basis for the award of any previous degree to "
                  f"anyone else.")
    for _ in range(5):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    p.paragraph_format.space_after = Pt(0)
    p.add_run("Date: ______________")
    p.add_run("\t\t\t\t")
    p.add_run("__________________________")
    for line in [GUIDE, "Designation: ____________________",
                 "Department of Computer Science and Engineering",
                 "Birla Institute of Technology, Mesra, Ranchi"]:
        q = doc.add_paragraph()
        q.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        q.paragraph_format.space_after = Pt(0)
        q.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        r = q.add_run(line); r.font.name = BODY_FONT; r.font.size = Pt(11.5)
    doc.add_paragraph()
    add_placeholder(doc, "The guide's designation has deliberately been left blank rather than "
                         "guessed. This page must be signed by the guide and must not be treated "
                         "as approved until it is.")


def build_declaration(doc):
    doc.add_page_break()
    centred(doc, "DECLARATION CERTIFICATE", size=14, bold=True, space_after=22)
    body_par(doc, "I certify that")
    items = [
        "The work contained in the thesis is original and has been done by myself under the general supervision of my supervisor.",
        "The work has not been submitted to any other Institute for any other degree or diploma.",
        "I have followed the guidelines provided by the Institute in writing the thesis.",
        "I have conformed to the norms and guidelines given in the Ethical Code of Conduct of the Institute.",
        "Whenever I have used materials (data, theoretical analysis, and text) from other sources, I have given due credit to them by citing them in the text of the thesis and giving their details in the references.",
        "Whenever I have quoted written materials from other sources, I have put them under quotation marks and given due credit to the sources by citing them and giving required details in the references.",
    ]
    for lab, text in zip("abcdef", items):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.left_indent = Inches(0.5)
        p.paragraph_format.first_line_indent = Inches(-0.3)
        p.paragraph_format.space_after = Pt(8)
        add_rich(p, f"{lab})\t{text}")
    for _ in range(4):
        doc.add_paragraph()
    for line in ["__________________________", STUDENT, f"({ROLL})", "Date: ______________"]:
        q = doc.add_paragraph()
        q.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        q.paragraph_format.space_after = Pt(0)
        q.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        q.add_run(line)
    doc.add_paragraph()
    add_placeholder(doc, "This declaration must be signed only after you have read the complete "
                         "thesis and satisfied yourself that every statement in it is true. It is "
                         "deliberately left unsigned here.")


def build_certificate(doc):
    doc.add_page_break()
    centred(doc, "CERTIFICATE OF APPROVAL", size=14, bold=True, space_after=24)
    body_par(doc, "This is to certify that the work embodied in this thesis entitled")
    body_par(doc, f"“{TITLE}”", align=WD_ALIGN_PARAGRAPH.CENTER)
    body_par(doc, f"is carried out by **{STUDENT}** ({ROLL}) and has been approved for the degree "
                  f"of **{DEGREE} in {BRANCH}** of Birla Institute of Technology, Mesra, Ranchi.")
    doc.add_paragraph()
    for line in ["Date: ______________", "Place: ______________"]:
        q = doc.add_paragraph()
        q.paragraph_format.space_after = Pt(2)
        q.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        q.add_run(line)
    for _ in range(5):
        doc.add_paragraph()
    t = doc.add_table(rows=1, cols=2)
    t.autofit = True
    for cell, label in zip(t.rows[0].cells, ["__________________\nInternal Examiner",
                                             "__________________\nExternal Examiner"]):
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        p.add_run(label)
    for _ in range(4):
        doc.add_paragraph()
    for line in ["__________________", "(Chairman)", "Head of the Department"]:
        q = doc.add_paragraph()
        q.alignment = WD_ALIGN_PARAGRAPH.CENTER
        q.paragraph_format.space_after = Pt(0)
        q.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        q.add_run(line)


def build_frontmatter_text(doc):
    text = (SRC / "00_frontmatter.md").read_text(encoding="utf-8")
    blocks = {}
    current = None
    for line in text.split("\n"):
        if line.startswith("@"):
            current = line[1:].strip()
            blocks[current] = []
        elif current:
            blocks[current].append(line)

    # Abstract
    doc.add_page_break()
    centred(doc, "ABSTRACT", size=14, bold=True, space_after=20)
    for line in blocks["ABSTRACT"]:
        if line.strip():
            body_par(doc, line.strip())
    kw = [l for l in blocks["KEYWORDS"] if l.strip()]
    if kw:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_before = Pt(10)
        r = p.add_run("Keywords: ")
        r.bold = True; r.font.name = BODY_FONT; r.font.size = Pt(12)
        add_rich(p, kw[0].strip(), size=Pt(12))

    # Acknowledgement
    doc.add_page_break()
    centred(doc, "ACKNOWLEDGEMENT", size=14, bold=True, space_after=20)
    for line in blocks["ACKNOWLEDGEMENT"]:
        s = line.strip()
        if not s:
            continue
        if s.startswith("!!"):
            add_placeholder(doc, s[2:].strip())
        else:
            body_par(doc, s)
    doc.add_paragraph()
    for line in ["Date: ______________"] + [l.strip() for l in blocks["SIGNOFF"] if l.strip()]:
        q = doc.add_paragraph()
        q.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        q.paragraph_format.space_after = Pt(0)
        q.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        q.add_run(line)


def build_contents(doc, cache=None):
    cache = cache or {}
    lvl = lambda n: "Normal"
    doc.add_page_break()
    centred(doc, "CONTENTS", size=14, bold=True, space_after=16)
    add_cached_toc(doc, r' TOC \o "1-3" \h \z \u ', cache.get("toc", []), lvl)

    doc.add_page_break()
    centred(doc, "LIST OF FIGURES", size=14, bold=True, space_after=16)
    add_cached_toc(doc, r' TOC \h \z \t "FigCaption,1" ', cache.get("lof", []), lvl)

    doc.add_page_break()
    centred(doc, "LIST OF TABLES", size=14, bold=True, space_after=16)
    add_cached_toc(doc, r' TOC \h \z \t "TblCaption,1" ', cache.get("lot", []), lvl)


# ----------------------------------------------------------------- main

def main(out_path, cache=None):
    doc = Document()
    configure_styles(doc)

    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.left_margin = Inches(1.5)
    sec.right_margin = Inches(1.0)
    sec.top_margin = Inches(1.0)
    sec.bottom_margin = Inches(1.0)
    section_page_numbering(sec, "lowerRoman", start=1)
    footer_page_number(sec)

    build_title_page(doc)
    build_approval(doc)
    build_declaration(doc)
    build_certificate(doc)
    build_frontmatter_text(doc)
    build_contents(doc, cache)

    # --- body section with arabic numbering restarting at 1 ---
    body = doc.add_section(WD_SECTION.NEW_PAGE)
    body.page_width = Cm(21.0)
    body.page_height = Cm(29.7)
    body.left_margin = Inches(1.5)
    body.right_margin = Inches(1.0)
    body.top_margin = Inches(1.0)
    body.bottom_margin = Inches(1.0)
    section_page_numbering(body, "decimal", start=1)
    footer_page_number(body)

    fig_list, tbl_list, outline = [], [], []
    state = {}
    order = ["01_intro.md", "02_background.md", "03_requirements.md",
             "04_architecture.md", "05_implementation.md", "06_results.md",
             "07_discussion.md", "08_conclusion.md", "09_appendices.md",
             "10_references.md"]
    first = True
    for name in order:
        path = SRC / name
        if first:
            # suppress the leading page break for the very first chapter
            text = path.read_text(encoding="utf-8")
            tmp = SRC / "._tmp.md"
            tmp.write_text(text, encoding="utf-8")
            render_first(doc, tmp, fig_list, tbl_list, outline, state)
            tmp.unlink()
            first = False
        else:
            render_source(doc, path, fig_list, tbl_list, outline, state)

    enable_update_fields(doc)
    doc.save(out_path)
    return fig_list, tbl_list, outline


def render_first(doc, path, fig_list, tbl_list, outline=None, state=None):
    """Render the first chapter without a page break before its heading."""
    state = {} if state is None else state
    state["first_chapter"] = True
    render_source(doc, path, fig_list, tbl_list, outline, state)


if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else str(HERE / "Tessera_Thesis.docx")
    figs, tbls, outline = main(out)
    print(f"written: {out}")
    print(f"figures: {len(figs)}  -> {[f[0] for f in figs]}")
    print(f"tables : {len(tbls)}  -> {[t[0] for t in tbls]}")
